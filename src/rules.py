#!/usr/bin/env python3
"""
五类数据质量规则校验

规则类型：
  完整性 - 必填字段不能为空
  唯一性 - 不能重复
  有效性 - 格式符合规范
  一致性 - 跨字段逻辑一致
  及时性 - 时间字段合理

用法:
  python rules.py --input samples/sample_customers.csv --report quality_report.xlsx
"""
import argparse
import os
import re
import sys
from datetime import datetime

import pandas as pd


class Rule:
    """规则基类"""

    def __init__(self, code, name, category, severity=1):
        self.code = code
        self.name = name
        self.category = category
        self.severity = severity

    def check(self, df):
        """返回 [(行号, 字段, 值, 描述)]"""
        raise NotImplementedError


class RequiredRule(Rule):
    """完整性：必填字段"""

    def __init__(self, field, code='R001'):
        super().__init__(code, f'{field} 必填', '完整性', 2)
        self.field = field

    def check(self, df):
        if self.field not in df.columns:
            return []
        mask = df[self.field].isna() | (df[self.field].astype(str).str.strip() == '')
        return [
            (idx + 2, self.field, '', f'{self.field} 为空')
            for idx in df[mask].index
        ]


class UniqueRule(Rule):
    """唯一性：字段组合不能重复"""

    def __init__(self, fields, code='R002'):
        super().__init__(code, f'{"+".join(fields)} 唯一', '唯一性', 2)
        self.fields = fields

    def check(self, df):
        if not all(f in df.columns for f in self.fields):
            return []
        dup_mask = df.duplicated(subset=self.fields, keep=False)
        result = []
        for idx in df[dup_mask].index:
            val = ' | '.join(str(df.loc[idx, f]) for f in self.fields)
            result.append((idx + 2, '+'.join(self.fields), val,
                           f'重复记录: {val}'))
        return result


class FormatRule(Rule):
    """有效性：格式校验"""

    def __init__(self, field, pattern, desc, code='R003'):
        super().__init__(code, f'{field} 格式', '有效性', 1)
        self.field = field
        self.pattern = re.compile(pattern)
        self.desc = desc

    def check(self, df):
        if self.field not in df.columns:
            return []
        result = []
        for idx, val in df[self.field].items():
            if pd.isna(val) or str(val).strip() == '':
                continue
            if not self.pattern.match(str(val).strip()):
                result.append((idx + 2, self.field, str(val), self.desc))
        return result


class ConsistencyRule(Rule):
    """一致性：跨字段逻辑"""

    def __init__(self, code='R004'):
        super().__init__(code, '时间字段逻辑一致', '一致性', 1)

    def check(self, df):
        """更新时间不能早于创建时间"""
        if 'created_at' not in df.columns or 'updated_at' not in df.columns:
            return []

        result = []
        for idx, row in df.iterrows():
            try:
                c = pd.to_datetime(row['created_at'])
                u = pd.to_datetime(row['updated_at'])
                if pd.notna(c) and pd.notna(u) and u < c:
                    result.append((idx + 2, 'updated_at', str(row['updated_at']),
                                   f'更新时间 {u} 早于创建时间 {c}'))
            except (ValueError, TypeError):
                continue
        return result


class TimelinessRule(Rule):
    """及时性：数据不能太旧"""

    def __init__(self, field, max_days=365, code='R005'):
        super().__init__(code, f'{field} 及时性', '及时性', 1)
        self.field = field
        self.max_days = max_days

    def check(self, df):
        if self.field not in df.columns:
            return []

        now = datetime.now()
        result = []
        for idx, val in df[self.field].items():
            if pd.isna(val):
                continue
            try:
                d = pd.to_datetime(val)
                days = (now - d).days
                if days > self.max_days:
                    result.append((idx + 2, self.field, str(val),
                                   f'已 {days} 天未更新，超过 {self.max_days} 天'))
            except (ValueError, TypeError):
                continue
        return result


def build_rules():
    """构建规则集"""
    return [
        RequiredRule('customer_name', 'R001'),
        RequiredRule('phone', 'R002'),
        UniqueRule(['phone', 'customer_name'], 'R003'),
        FormatRule('phone', r'^1[3-9]\d{9}$', '手机号应为 11 位、1 开头', 'R004'),
        FormatRule('tax_no', r'^[0-9A-Z]{15,20}$', '税号应为 15-20 位大写字母或数字', 'R005'),
        ConsistencyRule('R006'),
        TimelinessRule('updated_at', 365, 'R007'),
    ]


def main():
    p = argparse.ArgumentParser(description='数据质量规则校验')
    p.add_argument('--input', required=True)
    p.add_argument('--report', default='quality_report.xlsx')
    p.add_argument('--csv', default=None, help='同时导出 CSV')
    args = p.parse_args()

    if not os.path.exists(args.input):
        print(f'[!] 文件不存在: {args.input}')
        sys.exit(1)

    print(f'[*] 读取 {args.input} ...')
    df = pd.read_csv(args.input, dtype=str)
    print(f'[*] {len(df):,} 行')
    print()

    rules = build_rules()
    all_issues = []

    print('=' * 90)
    print('  规则校验')
    print('=' * 90)

    for rule in rules:
        issues = rule.check(df)
        status = 'PASS' if not issues else f'{len(issues):>5} 处'
        print(f'  {rule.code} [{rule.category}] {rule.name:<24} {status}')

        for row_no, field, val, desc in issues:
            all_issues.append({
                '行号': row_no,
                '规则编号': rule.code,
                '规则类别': rule.category,
                '规则名称': rule.name,
                '字段': field,
                '字段值': str(val)[:100],
                '问题描述': desc,
                '严重程度': '错误' if rule.severity == 2 else '警告',
            })

    print()
    print('=' * 90)
    print(f'  合计 {len(all_issues):,} 个问题')
    print('=' * 90)
    print()

    if not all_issues:
        print('[+] 未发现问题')
        return

    # 按规则汇总
    df_issues = pd.DataFrame(all_issues)
    summary = df_issues.groupby(['规则编号', '规则名称', '严重程度']).size().reset_index(name='数量')
    print('  按规则汇总：')
    for _, r in summary.iterrows():
        print(f'    {r["规则编号"]} {r["规则名称"]:<24} {r["数量"]:>6} 处  [{r["严重程度"]}]')
    print()

    # 导出 Excel
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment

        wb = Workbook()

        # Sheet1: 问题清单
        ws = wb.active
        ws.title = '问题清单'

        headers = list(df_issues.columns)
        ws.append(headers)

        for cell in ws[1]:
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='1F4E79')
            cell.alignment = Alignment(horizontal='center')

        err_fill = PatternFill('solid', fgColor='FFC7CE')
        warn_fill = PatternFill('solid', fgColor='FFEB9C')

        for _, r in df_issues.iterrows():
            ws.append([r[h] for h in headers])
            row = ws.max_row
            fill = err_fill if r['严重程度'] == '错误' else warn_fill
            for col in range(1, len(headers) + 1):
                ws.cell(row, col).fill = fill

        widths = [8, 10, 12, 22, 16, 30, 40, 10]
        for i, w in enumerate(widths[:len(headers)], 1):
            ws.column_dimensions[chr(64 + i)].width = w
        ws.freeze_panes = 'A2'
        ws.auto_filter.ref = f'A1:{chr(64 + len(headers))}{ws.max_row}'

        # Sheet2: 汇总
        ws2 = wb.create_sheet('汇总')
        ws2.append(['规则编号', '规则名称', '严重程度', '数量'])
        for cell in ws2[1]:
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='1F4E79')
            cell.alignment = Alignment(horizontal='center')

        for _, r in summary.iterrows():
            ws2.append([r['规则编号'], r['规则名称'], r['严重程度'], r['数量']])

        for i, w in enumerate([10, 24, 10, 10], 1):
            ws2.column_dimensions[chr(64 + i)].width = w

        wb.save(args.report)
        print(f'[+] Excel 报告: {args.report}')

    except ImportError:
        print('[!] 需要 openpyxl: pip install openpyxl')

    if args.csv:
        df_issues.to_csv(args.csv, index=False, encoding='utf-8-sig')
        print(f'[+] CSV: {args.csv}')

    print()
    print('  下一步：')
    print('    python src/clean_phone.py --input', args.input)


if __name__ == '__main__':
    main()
